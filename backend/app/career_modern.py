"""Versioned single-column layout, shared by DOCX and PDF exports."""
from io import BytesIO
from pathlib import Path
import re
from html import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, HRFlowable


TOKENS = {
    "version": "career_modern.1", "font": "Arial", "fallback_font": "Helvetica",
    "margin": .7, "body": 10.5, "name": 23, "section": 11, "role": 11,
    "contact": 9, "leading": 1.16, "after": 5, "section_before": 12,
    "bullet_indent": 12, "bullet_hanging": 8, "rule_width": .5,
    "ink": "17243A", "accent": "087E8B", "organisation": "6D4288", "muted": "596674",
    "cover_ink": "111111",
}


def export_content(content: str) -> str:
    """Remove empty sections and their whitespace, preserving all body facts."""
    lines = content.strip().splitlines()
    kept = []
    for index, line in enumerate(lines):
        heading = re.match(r"^(#{1,3})\s", line.strip())
        if heading:
            has_content = False
            for value in lines[index + 1:]:
                value = value.strip()
                next_heading = re.match(r"^(#{1,3})\s", value)
                if next_heading and len(next_heading[1]) <= len(heading[1]):
                    break
                if value and not next_heading:
                    has_content = True
                    break
            if not has_content:
                continue
        kept.append(line)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(kept)).strip()


def blocks(content: str):
    lines = [line.strip() for line in export_content(content).splitlines() if line.strip()]
    while len(lines) > 2 and re.match(r"(?i)^(?:Phone|Email|Location):", lines[1]) and re.match(r"(?i)^(?:Phone|Email|Location):", lines[2]):
        lines[1:3] = [lines[1] + " | " + lines[2]]
    previous = ""
    for index, line in enumerate(lines):
        kind = "body"
        if index == 0:
            kind = "name"
        elif index == 1 and not line.startswith("#") and ("@" in line or "|" in line):
            kind = "contact"
        elif line.startswith("### "):
            kind = "role"
        elif line.startswith("## ") or line.startswith("# "):
            kind = "section"
        elif line.startswith("Cover Letter:"):
            kind = "role"
        elif line.startswith("- "):
            kind = "bullet"
        elif previous == "role":
            kind = "organisation"
        text = re.sub(r"^#{1,3}\s+|^-\s+", "", line)
        if kind == "section":
            text = text.upper()
        period = re.search(r"(?i)(?:\b(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\s+)?\b(?:19|20)\d{2}\s*-\s*(?:(?:[A-Za-z]+\s+)?(?:19|20)\d{2}|Present|Current|Ongoing)\s*$", text) if kind == "organisation" else None
        if period:
            if text[:period.start()].strip(" |"):
                yield kind, text[:period.start()].strip(" |")
            yield "date", period[0]
        else:
            yield kind, text
        previous = kind


def create_modern(content: str, title: str, format: str) -> bytes:
    from .exporter import _add_bottom_border, _ascii_punctuation
    tokens = {**TOKENS, "ink": TOKENS["cover_ink"], "organisation": TOKENS["muted"]} if title.casefold() == "cover letter" else TOKENS
    stream = BytesIO()
    items = list(blocks(_ascii_punctuation(content)))
    if format == "docx":
        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Pt(A4[0]), Pt(A4[1])
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(tokens["margin"])
        normal = doc.styles["Normal"]
        normal.font.name = tokens["font"]
        normal.font.size = Pt(tokens["body"])
        for kind, text in items:
            p = doc.add_paragraph()
            f = p.paragraph_format
            f.space_after = Pt(tokens["after"])
            f.line_spacing = tokens["leading"]
            f.widow_control = True
            f.keep_with_next = kind in {"name", "contact", "section", "role", "organisation", "date"}
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if kind == "date" else WD_ALIGN_PARAGRAPH.LEFT
            f.space_before = Pt(tokens["section_before"] if kind == "section" else 0)
            if kind == "bullet":
                f.left_indent = Pt(tokens["bullet_indent"])
                f.first_line_indent = Pt(-tokens["bullet_hanging"])
                text = "• " + text
            for fragment in re.split(r"(\*\*.*?\*\*)", text):
                run = p.add_run(fragment.strip("*") if fragment.startswith("**") else fragment)
                run.font.name = tokens["font"]
                run.font.size = Pt(tokens.get(kind, tokens["body"]) if kind in {"name", "section", "role", "contact"} else tokens["body"])
                run.bold = kind in {"name", "section", "role"} or fragment.startswith("**")
                run.font.color.rgb = RGBColor.from_string(tokens["accent"] if kind == "section" else tokens["organisation"] if kind == "organisation" else tokens["muted"] if kind == "contact" else tokens["ink"])
            if kind in {"contact", "section"}:
                _add_bottom_border(p, RGBColor.from_string(tokens["accent"] if kind == "section" else tokens["muted"]), tokens["rule_width"])
        doc.core_properties.comments = "Template tokens: " + tokens["version"]
        doc.save(stream)
    else:
        # Explicit Arial font embedding where available; built-in Helvetica fallback
        # keeps the same point sizes. No network font download is needed.
        font, bold = "Helvetica", "Helvetica-Bold"
        font_path = Path("C:/Windows/Fonts/arial.ttf")
        bold_path = Path("C:/Windows/Fonts/arialbd.ttf")
        if font_path.exists() and bold_path.exists():
            if "CareerArial" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("CareerArial", str(font_path)))
                pdfmetrics.registerFont(TTFont("CareerArial-Bold", str(bold_path)))
                pdfmetrics.registerFontFamily("CareerArial", normal="CareerArial", bold="CareerArial-Bold", italic="CareerArial", boldItalic="CareerArial-Bold")
            font, bold = "CareerArial", "CareerArial-Bold"
        story = []
        for kind, text in items:
            size = tokens[kind] if kind in {"name", "section", "role", "contact"} else tokens["body"]
            style = ParagraphStyle(kind, fontName=bold if kind in {"name", "section", "role"} else font,
                                   fontSize=size, leading=size * tokens["leading"], spaceAfter=tokens["after"],
                                   spaceBefore=tokens["section_before"] if kind == "section" else 0,
                                   textColor="#" + tokens["accent" if kind == "section" else "organisation" if kind == "organisation" else "muted" if kind == "contact" else "ink"],
                                   keepWithNext=kind in {"name", "contact", "section", "role", "organisation", "date"},
                                   alignment=TA_RIGHT if kind == "date" else TA_LEFT,
                                   leftIndent=tokens["bullet_indent"] if kind == "bullet" else 0,
                                   bulletIndent=tokens["bullet_indent"] - tokens["bullet_hanging"])
            rendered = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", escape(text))
            story.append(Paragraph(rendered, style, bulletText="•" if kind == "bullet" else None))
            if kind in {"contact", "section"}:
                rule = HRFlowable(width="100%", thickness=tokens["rule_width"], color="#" + tokens["accent" if kind == "section" else "muted"], spaceAfter=tokens["after"])
                rule.keepWithNext = True
                story.append(rule)
        margin = tokens["margin"] * 72
        SimpleDocTemplate(stream, pagesize=A4, rightMargin=margin, leftMargin=margin,
                          topMargin=margin, bottomMargin=margin, title=title, subject=tokens["version"]).build(story)
    return stream.getvalue()
