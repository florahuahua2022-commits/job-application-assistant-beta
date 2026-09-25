import json
import unittest
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine, select

from app.application_decision import decision_inputs
from app.ai import AIServiceError
from app.application_requirements import empty_application_requirements
from app.ckb import build_career_knowledge_base
from app.database import get_session
from app.main import app, get_or_refresh_current_ckb, master_resume_integrity_issue, repair_legacy_resume_evidence
from app.models import ApplicantProfile, GeneratedDocument, JobApplication, JobSource, Resume
from app.outcome_learning import build_submission_snapshot
from app.release_state import details_fingerprint, fingerprint, generation_inputs_fingerprint, pack_fingerprint
from app.resume_plan import build_resume_curation_plan, validate_resume_content
from app.ingest import parse_job_ad_text


PRODUCTION_MISSING_EXPERIENCES_SOURCE = """WORK EXPERIENCE
Department of Communities - Disability Services, WA State Government February 2026 - August 2026
Finance Administration Officer
Provided administrative and operational support to business units.
Self-employed via Mable August 2025 - January 2026
Independent Support Worker
Managed client scheduling, appointments and direct communication independently, maintaining professional service records.
Provided individual support with daily living and community participation according to each client's needs.
My Support August 2024 - August 2025
Support Worker
Delivered individualised support with daily living, appointments and community access.
Additional Australian employment: October 2022 - April 2024
Service and casual roles
Sodex: Utility, December 2023 - April 2024.
Woolworths: Cashier, May 2023 - November 2023.
Puma, Port Hedland: service station work, March 2023 - May 2023.
SA Health: casual engagement, October 2022 (one month).
Core Color, Adelaide May 2022 - September 2022
E-commerce Operations
Processed supplier orders through a CRM system.
Self-employed - Amazon e-commerce business 2019 - 2022
Self-employed E-commerce Operator
Operated an independent Amazon e-commerce business.
EDUCATION
Bachelor of Arts"""

PRODUCTION_SAVED_EXPERIENCES = json.dumps([{
    "id": "finance", "role_title": "Finance Administration Officer",
    "organization": "Department of Communities - Disability Services, WA State Government",
    "time_period_text": "February 2026 - August 2026",
    "responsibility": "User-edited wording.",
}])

PRODUCTION_CLEAN_NINE_EXPERIENCES = [{
    "id": "EVA37A578823CE", "role_title": "Finance Administration Officer",
    "organization": "Department of Communities - Disability Services, WA State Government",
    "time_period_text": "February 2026 - August 2026", "responsibility": "Provided administrative support.",
}, {
    "id": "EV2572ECCB036E",
    "role_title": "Assisted in prioritising competing tasks to support service delivery; used Dayforce within a WA Government environment. Processed journals and completed reconciliations.",
    "organization": "Self-employed via Mable", "time_period_text": "August 2025 - January 2026", "responsibility": "",
}, {
    "id": "EV849245D0F4F5",
    "role_title": "Provided individual support to independently sourced clients through Mable, including community participation, appointments and assistance with daily living.",
    "organization": "My Support", "time_period_text": "August 2024 - August 2025", "responsibility": "",
}, {
    "id": "woolworths", "role_title": "Cashier", "organization": "Woolworths",
    "time_period_text": "May 2023 - November 2023", "responsibility": "Casual cashier work.",
}, {
    "id": "core-color", "role_title": "E-commerce Operations", "organization": "Core Color",
    "time_period_text": "May 2022 - September 2022", "responsibility": "Processed supplier orders through a CRM system.",
}, {
    "id": "EV2FAF7A3012FA", "role_title": "Executive Assistant to Board Member", "organization": "Avaintec",
    "time_period_text": "November 2017 - January 2019", "responsibility": "Administrative support.",
}, {
    "id": "EV51BD90C637AE", "role_title": "Project Administration Officer",
    "organization": "China Communications Construction Company - Kenya Branch",
    "time_period_text": "January 2016 - August 2017", "responsibility": "Project support.",
}, {
    "id": "EV1BFA570A50FD", "role_title": "Project Administration Officer", "organization": "Chevron CDB Project",
    "time_period_text": "August 2012 - December 2015", "responsibility": "Project support.",
}, {
    "id": "EV8F624C6303E4", "role_title": "Project Assistant", "organization": "Pratt & Whitney",
    "time_period_text": "October 2007 - August 2012", "responsibility": "Project support.",
}]


class RealUserRegressionTests(unittest.TestCase):
    def test_exclusion_api_is_validated_idempotent_and_reversible(self):
        source = """WORK EXPERIENCE
Example Agency January 2020 - Present
Project Officer
Prepared reports and coordinated meetings.
Sodex: Utility, December 2023 - April 2024.
EDUCATION"""
        experiences = json.dumps([{
            "id": "saved", "role_title": "Project Officer", "organization": "Example Agency",
            "time_period_text": "January 2020 - Present", "responsibility": "Prepared reports and coordinated meetings.",
        }])
        with Session(self.engine) as session:
            resume = Resume(title="Master Resume", source_text=source, experiences_json=experiences, ckb_json='[{"sentinel":true}]')
            session.add(resume); session.commit(); session.refresh(resume); resume_id = resume.id

        scan = self.client.post("/resumes/risk-scan")
        candidate = scan.json()["resumes"][0]["experiences"][0]
        self.assertEqual(candidate["status"], "unresolved")
        before = self.client.get("/resumes").json()[0]

        rejected = self.client.patch(f"/resumes/{resume_id}/experience-exclusions", json={"candidate_id": "EXUNKNOWN", "action": "exclude"})
        self.assertEqual(rejected.status_code, 422, rejected.text)

        payload = {"candidate_id": candidate["candidate_id"], "action": "exclude"}
        first = self.client.patch(f"/resumes/{resume_id}/experience-exclusions", json=payload)
        second = self.client.patch(f"/resumes/{resume_id}/experience-exclusions", json=payload)
        self.assertEqual(first.status_code, 200, first.text)
        self.assertEqual(second.status_code, 200, second.text)
        self.assertEqual(first.json()["exclusions"], second.json()["exclusions"])
        self.assertEqual(first.json()["candidates"][0]["status"], "excluded_by_user")

        after = self.client.get("/resumes").json()[0]
        self.assertEqual(after["source_text"], before["source_text"])
        self.assertEqual(after["experiences_json"], before["experiences_json"])
        self.assertEqual(after["ckb_json"], before["ckb_json"])
        self.assertEqual(self.client.post("/resumes/risk-scan").json()["needs_review_count"], 0)

        restored = self.client.patch(f"/resumes/{resume_id}/experience-exclusions", json={**payload, "action": "restore"})
        self.assertEqual(restored.status_code, 200, restored.text)
        self.assertEqual(restored.json()["exclusions"], [])
        self.assertEqual(restored.json()["candidates"][0]["status"], "unresolved")

    def test_real_sodex_puma_amazon_candidates_are_excluded_from_rebuilt_ckb(self):
        source = """WORK EXPERIENCE
Example Agency January 2020 - Present
Project Officer
Prepared reports and coordinated meetings.
Sodex: Utility, December 2023 - April 2024.
Puma, Port Hedland: service station work, March 2023 - May 2023.
Self-employed - Amazon e-commerce business 2019 - 2022
Self-employed E-commerce Operator
Operated an independent Amazon e-commerce business.
EDUCATION"""
        experiences = json.dumps([{
            "id": "saved", "role_title": "Project Officer", "organization": "Example Agency",
            "time_period_text": "January 2020 - Present", "responsibility": "Prepared reports and coordinated meetings.",
        }])
        with Session(self.engine) as session:
            resume = Resume(title="Master Resume", source_text=source, experiences_json=experiences, ckb_json='[{"sentinel":true}]')
            session.add(resume); session.commit(); session.refresh(resume); resume_id = resume.id

        scan = self.client.post("/resumes/risk-scan").json()
        candidates = scan["resumes"][0]["experiences"]
        self.assertEqual(
            {(item["organization"], item["role_title"], item["time_period_text"]) for item in candidates},
            {
                ("Sodex", "Utility", "December 2023 - April 2024"),
                ("Puma, Port Hedland", "service station work", "March 2023 - May 2023"),
                ("Self-employed - Amazon e-commerce business", "Self-employed E-commerce Operator", "2019 - 2022"),
            },
        )
        for item in candidates:
            response = self.client.patch(f"/resumes/{resume_id}/experience-exclusions", json={
                "candidate_id": item["candidate_id"], "action": "exclude",
            })
            self.assertEqual(response.status_code, 200, response.text)

        self.assertEqual(self.client.post("/resumes/risk-scan").json()["needs_review_count"], 0)
        with Session(self.engine) as session:
            stored = session.get(Resume, resume_id)
            excluded = stored.experience_exclusions_json
            candidate_experiences = json.dumps(json.loads(experiences) + [{
                "role_title": "Utility", "organization": "Sodex",
                "time_period_text": "December 2023 - April 2024", "responsibility": "Casual work.",
            }])
            rebuilt = build_career_knowledge_base(source, candidate_experiences, excluded)
            self.assertEqual([item["source_section"] for item in rebuilt], ["Work Experience > Example Agency > Project Officer"])
            plan = build_resume_curation_plan(
                {"criteria": [{"criteria_id": "C1", "criteria_type": "essential"}]},
                {"matches": [{"criteria_id": "C1", "matched_evidence": ["excluded-sodex"], "coverage": "strong"}]},
                rebuilt,
            )
            self.assertNotIn("excluded-sodex", plan["selected_evidence"])
            self.assertEqual(stored.source_text, source)
            self.assertIsNone(master_resume_integrity_issue(stored))

    def test_legacy_evidence_repair_refreshes_existing_snapshots_and_caches_idempotently(self):
        legacy = json.dumps([{
            "id": "EVD886D211EFA3", "role_title": "Additional Australian experience",
            "organization": "Sodex: Utility, . Woolworths: Cashier, May 2023 - November 2023.",
            "time_period_text": "December 2023 - April 2024", "responsibility": "",
        }, {
            "id": "EV1CBFDD076607",
            "role_title": "Sodex: Utility, December 2023 - April 2024. Woolworths: Cashier, May 2023 - November 2023.",
            "organization": "Puma, Port Hedland: service station work, .",
            "time_period_text": "March 2023 - May 2023",
            "responsibility": "Core Color 2022 E-commerce operations Adelaide Processed supplier orders through a CRM system.",
        }, {
            "id": "older-role", "role_title": "Executive Assistant", "organization": "Older Employer",
            "time_period_text": "2017 - 2019", "responsibility": "Administrative support.",
        }])
        exclusions = json.dumps([{
            "organization": "Sodex", "role_title": "Utility", "time_period_text": "December 2023 - April 2024",
        }, {
            "organization": "Puma, Port Hedland", "role_title": "service station work", "time_period_text": "March 2023 - May 2023",
        }, {
            "organization": "Self-employed - Amazon e-commerce business", "role_title": "Self-employed e-commerce operator",
            "time_period_text": "2019 - 2022",
        }])
        with Session(self.engine) as session:
            resume = Resume(title="Master Resume", source_text=PRODUCTION_MISSING_EXPERIENCES_SOURCE,
                            experiences_json=legacy, ckb_json=json.dumps(build_career_knowledge_base(PRODUCTION_MISSING_EXPERIENCES_SOURCE, legacy)),
                            experience_exclusions_json=exclusions)
            session.add(resume); session.commit(); session.refresh(resume)
            stale_snapshot = json.dumps({
                "resume_id": resume.id, "source_text": resume.source_text, "experiences_json": legacy,
                "ckb_json": resume.ckb_json, "experience_exclusions_json": exclusions,
            })
            curtin = JobApplication(company="Curtin University", position_title="Fieldwork Administrative Support Officer",
                                    job_description="Administration", resume_snapshot_json=stale_snapshot,
                                    evidence_matches_json='{"matches":["EV1CBFDD076607"]}', selection_plan_json='{"items":[1]}')
            sra = JobApplication(company="SRA Solutions", position_title="Project Administrator / Document Controller",
                                 job_description="Project administration", resume_snapshot_json=stale_snapshot,
                                 evidence_matches_json='{"matches":["EV1CBFDD076607"]}', selection_plan_json='{"items":[1]}')
            session.add_all([curtin, sra]); session.commit()

            first = repair_legacy_resume_evidence(session, resume, None)
            first_state = (resume.experiences_json, resume.ckb_json, curtin.resume_snapshot_json, sra.resume_snapshot_json)
            second = repair_legacy_resume_evidence(session, resume, None)
            second_state = (resume.experiences_json, resume.ckb_json, curtin.resume_snapshot_json, sra.resume_snapshot_json)

        self.assertEqual((first, second), (2, 0))
        self.assertEqual(first_state, second_state)
        cleaned = json.loads(resume.experiences_json)
        self.assertEqual([item["organization"] for item in cleaned], ["Woolworths", "Core Color", "Older Employer"])
        self.assertNotIn("EV1CBFDD076607", resume.ckb_json)
        for application in (curtin, sra):
            snapshot = json.loads(application.resume_snapshot_json)
            self.assertEqual(snapshot["experiences_json"], resume.experiences_json)
            self.assertEqual(snapshot["ckb_json"], resume.ckb_json)
            self.assertEqual((application.evidence_matches_json, application.selection_plan_json), ("{}", "{}"))

    def test_excluding_after_application_snapshot_requires_update_to_latest(self):
        source = """WORK EXPERIENCE
Example Agency January 2020 - Present
Project Officer
Prepared reports and coordinated meetings.
Sodex: Utility, December 2023 - April 2024.
EDUCATION"""
        experiences = json.dumps([{
            "id": "saved", "role_title": "Project Officer", "organization": "Example Agency",
            "time_period_text": "January 2020 - Present", "responsibility": "Prepared reports and coordinated meetings.",
        }])
        old_snapshot = json.dumps({
            "resume_id": 1, "title": "Master Resume", "source_text": source,
            "experiences_json": experiences, "ckb_json": "[]", "experience_exclusions_json": "[]",
        })
        with Session(self.engine) as session:
            resume = Resume(title="Master Resume", source_text=source, experiences_json=experiences, ckb_json="[]")
            application = JobApplication(company="Curtin University", position_title="Fieldwork Administrative Support Officer",
                                         job_description="Provide administrative support.", resume_snapshot_json=old_snapshot)
            session.add_all([resume, application]); session.commit(); session.refresh(resume); session.refresh(application)
            resume_id, application_id = resume.id, application.id

        candidate = self.client.post("/resumes/risk-scan").json()["resumes"][0]["experiences"][0]
        excluded = self.client.patch(f"/resumes/{resume_id}/experience-exclusions", json={
            "candidate_id": candidate["candidate_id"], "action": "exclude",
        })
        self.assertEqual(excluded.status_code, 200, excluded.text)

        blocked = self.client.post("/generate", json={"application_id": application_id, "document_type": "tailored_resume"})
        self.assertEqual(blocked.status_code, 409, blocked.text)
        self.assertEqual(blocked.json()["detail"]["code"], "application_resume_snapshot_outdated")

        updated = self.client.put(f"/applications/{application_id}/resume", json={
            "use_latest_master": True, "source_text": None, "expected_snapshot": old_snapshot,
        })
        self.assertEqual(updated.status_code, 200, updated.text)
        snapshot = json.loads(updated.json()["resume_snapshot_json"])
        self.assertEqual(json.loads(snapshot["experience_exclusions_json"])[0]["candidate_id"], candidate["candidate_id"])

    def test_update_to_latest_rejects_resume_with_uncovered_experiences_without_replacing_snapshot(self):
        with Session(self.engine) as session:
            resume = Resume(title="Master Resume", source_text=PRODUCTION_MISSING_EXPERIENCES_SOURCE,
                            experiences_json=PRODUCTION_SAVED_EXPERIENCES, ckb_json="[]")
            application = JobApplication(
                company="Curtin University", position_title="Fieldwork Administrative Support Officer",
                job_description="Provide administrative support.", resume_snapshot_json='{"sentinel":"unchanged"}',
            )
            session.add_all([resume, application]); session.commit(); session.refresh(application)
            application_id = application.id

        response = self.client.put(f"/applications/{application_id}/resume", json={
            "use_latest_master": True, "source_text": None, "expected_snapshot": '{"sentinel":"unchanged"}',
        })

        self.assertEqual(response.status_code, 409, response.text)
        self.assertEqual(response.json()["detail"]["code"], "master_resume_experience_needs_review")
        self.assertIn("Self-employed via Mable", str(response.json()["detail"]))
        with Session(self.engine) as session:
            self.assertEqual(session.get(JobApplication, application_id).resume_snapshot_json, '{"sentinel":"unchanged"}')

    def test_missing_experience_coverage_saves_but_still_blocks_scan_and_generation(self):
        created = self.client.post("/resumes", json={
            "title": "Master Resume", "source_text": PRODUCTION_MISSING_EXPERIENCES_SOURCE,
            "experiences_json": PRODUCTION_SAVED_EXPERIENCES,
        })
        self.assertEqual(created.status_code, 200, created.text)
        self.assertGreater(created.json()["unresolved_experience_count"], 0)

        safe_source = "Work Experience\nProject Officer Example Agency Feb 2020 - Present\nProject Officer\nPrepared reports."
        with Session(self.engine) as session:
            resume = Resume(title="Master Resume", source_text=safe_source, experiences_json=json.dumps([{
                "id": "safe", "role_title": "Project Officer", "organization": "Example Agency",
                "time_period_text": "Feb 2020 - Present", "responsibility": "Prepared reports.",
            }]), ckb_json="[]")
            application = JobApplication(
                company="Example", position_title="Administrator", job_description="Administrative support.",
                job_model_json='{"schema_version":"1.0","criteria":[],"limit_scope":"unspecified"}',
                application_requirements_json=json.dumps(self.requirements("resume")),
            )
            session.add_all([resume, application]); session.commit(); session.refresh(resume); session.refresh(application)
            resume_id, application_id = resume.id, application.id

        before = self.client.get("/resumes").json()[0]
        edited = self.client.patch(f"/resumes/{resume_id}", json={
            "source_text": PRODUCTION_MISSING_EXPERIENCES_SOURCE,
            "experiences_json": PRODUCTION_SAVED_EXPERIENCES,
        })
        self.assertEqual(edited.status_code, 200, edited.text)
        self.assertGreater(edited.json()["unresolved_experience_count"], 0)

        with Session(self.engine) as session:
            resume = session.get(Resume, resume_id)
            resume.source_text = PRODUCTION_MISSING_EXPERIENCES_SOURCE
            resume.experiences_json = PRODUCTION_SAVED_EXPERIENCES
            session.add(resume); session.commit()

        scan = self.client.post("/resumes/risk-scan")
        self.assertEqual(scan.status_code, 200, scan.text)
        self.assertGreaterEqual(scan.json()["needs_review_count"], 1)
        self.assertIn("Core Color, Adelaide", str(scan.json()["resumes"]))
        with Session(self.engine) as session:
            self.assertEqual(session.get(Resume, resume_id).experiences_json, PRODUCTION_SAVED_EXPERIENCES)

        generated = self.client.post("/generate", json={
            "application_id": application_id, "document_type": "tailored_resume",
        })
        self.assertEqual(generated.status_code, 409, generated.text)
        self.assertEqual(generated.json()["detail"]["code"], "master_resume_experience_needs_review")

    def test_upload_saves_when_real_resume_experiences_are_omitted_by_parsing(self):
        parsed = json.loads(PRODUCTION_SAVED_EXPERIENCES)
        with patch("app.main.extract_resume_text", return_value=PRODUCTION_MISSING_EXPERIENCES_SOURCE), patch(
            "app.main.extract_resume_experiences", return_value=parsed,
        ):
            response = self.client.post(
                "/resumes/upload",
                files={"file": ("resume.docx", b"document", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={"title": "Master Resume"},
            )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertGreater(response.json()["unresolved_experience_count"], 0)

    def test_upload_unresolved_then_exclude_allows_update_and_generation(self):
        application_id = self.seed(required=("resume",))
        source = """WORK EXPERIENCE
Example Agency January 2020 - Present
Project Officer
Prepared reports and coordinated meetings.
Sodex: Utility, December 2023 - April 2024.
EDUCATION"""
        parsed = [{
            "id": "saved", "role_title": "Project Officer", "organization": "Example Agency",
            "time_period_text": "January 2020 - Present", "responsibility": "Prepared reports and coordinated meetings.",
        }]
        with patch("app.main.extract_resume_text", return_value=source), patch("app.main.extract_resume_experiences", return_value=parsed):
            uploaded = self.client.post("/resumes/upload", files={"file": ("resume.docx", b"document")}, data={"title": "Master Resume"})
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        self.assertEqual(uploaded.json()["unresolved_experience_count"], 1)

        with Session(self.engine) as session:
            application = session.get(JobApplication, application_id)
            old_snapshot = application.resume_snapshot_json
        blocked_update = self.client.put(f"/applications/{application_id}/resume", json={
            "use_latest_master": True, "source_text": None, "expected_snapshot": old_snapshot,
        })
        self.assertEqual(blocked_update.status_code, 409, blocked_update.text)

        issue = self.client.post("/resumes/risk-scan").json()["resumes"][0]["experiences"][0]
        excluded = self.client.patch(f"/resumes/{uploaded.json()['id']}/experience-exclusions", json={
            "candidate_id": issue["candidate_id"], "action": "exclude",
        })
        self.assertEqual(excluded.status_code, 200, excluded.text)
        updated = self.client.put(f"/applications/{application_id}/resume", json={
            "use_latest_master": True, "source_text": None, "expected_snapshot": old_snapshot,
        })
        self.assertEqual(updated.status_code, 200, updated.text)

        with Session(self.engine) as session:
            application = session.get(JobApplication, application_id)
            profile = session.exec(select(ApplicantProfile)).first()
            requirements = json.loads(application.application_requirements_json)
            application.application_decision_json = json.dumps({
                "schema_version": "1.0", "status": "ready", "application_recommendation": "apply",
                "inputs": decision_inputs(json.loads(application.job_model_json), requirements, [], profile),
                "requirements": [], "questions": [], "blocking_issues": [],
            })
            session.add(application); session.commit()
        plan = {"selected_evidence": [], "roles": [{
            "employer_marker": "Example Agency", "role_marker": "Project Officer",
            "display_period": "January 2020 - Present", "chronology_order": 0, "include_role_header": True,
        }]}
        draft = "## Professional Summary\nGrounded support.\n## Key Skills\nAdministration\n## Work Experience\n**Project Officer**\n**Example Agency**\nJanuary 2020 - Present"
        with patch("app.main.match_evidence_batch", return_value={"schema_version": "1.0", "matches": [], "unused_evidence": []}), patch(
            "app.main.build_resume_curation_plan", return_value=plan,
        ), patch("app.main.generate_draft", return_value=draft), patch(
            "app.main.repair_tailored_resume", return_value=(draft, {"status": "pass", "results": []}),
        ):
            generated = self.client.post("/generate", json={"application_id": application_id, "document_type": "tailored_resume"})
        self.assertEqual(generated.status_code, 200, generated.text)

    def test_uploaded_experiences_persist_review_annotations(self):
        source = "Work Experience\nProject Officer\nExample Agency\nFeb 2020 - Present\nPrepared reports.\nEducation"
        with patch("app.main.extract_resume_text", return_value=source):
            response = self.client.post(
                "/resumes/upload", files={"file": ("resume.docx", b"document", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={"title": "Master Resume"},
            )

        self.assertEqual(response.status_code, 200, response.text)
        with Session(self.engine) as session:
            experience = json.loads(session.exec(select(Resume)).first().experiences_json)[0]
        self.assertIs(experience["needs_review"], False)
        self.assertEqual(experience["review_reasons"], [])

    def test_bennco_pasted_ad_persists_only_formal_criteria_in_job_model(self):
        raw_text = """Project Administrator
Bennco Group
Selection Criteria
Experience in construction or mining (preferred)
Who We Are
Bennco Group is a multi-disciplinary building and construction contractor supporting Tier 1 clients across the Pilbara and wider WA.
Our Values
Pride & Commitment – We own our work and get the job done.
Growth & Improvement – We push ourselves to evolve and excel.
Family & Loyalty – We look after our people and create a welcoming team culture.
Trust & Respect – We communicate openly and honour our commitments.
"""
        parsed = parse_job_ad_text(raw_text)
        created = self.client.post("/applications", json={
            "company": "Bennco Group", "position_title": "Project Administrator",
            "job_description": parsed["job_description"], "selection_criteria": parsed["selection_criteria"],
        })

        self.assertEqual(created.status_code, 200, created.text)
        application = created.json()
        model = json.loads(application["job_model_json"])
        self.assertEqual(application["selection_criteria"], "Experience in construction or mining (preferred)")
        self.assertEqual(model["requirement_mode"], "inferred_requirements")
        self.assertEqual([item["criteria_text"] for item in model["criteria"]], ["Experience in construction or mining (preferred)"])

    def test_pasted_complete_resume_builds_canonical_employment_before_ckb(self):
        source = """Work Experience
Project Officer
Example Agency
Feb 2020 – Present
Prepared reports, coordinated meetings and maintained accurate project records.
"""
        response = self.client.post("/resumes", json={"title": "Master Resume", "source_text": source, "experiences_json": "[]"})
        self.assertEqual(response.status_code, 200, response.text)
        with Session(self.engine) as session:
            resume = session.exec(select(Resume)).first()
        experience = json.loads(resume.experiences_json)[0]
        evidence = json.loads(resume.ckb_json)[0]
        self.assertEqual(experience["time_period_text"], "Feb 2020 – Present")
        self.assertEqual(evidence["time_period"], {"start": "Feb 2020", "end": "Present"})

    def test_historical_resume_risk_scan_marks_without_rewriting_and_blocks_generation(self):
        application_id = self.seed(required=("resume",))
        source = """Work Experience
Support Worker
Provided individual support to clients through Mable, including appointments and assistance with daily living.
Mable | Jan 2020 – Present
Education"""
        experience = {
            "role_title": "Provided individual support to clients through Mable, including appointments and assistance with daily living.",
            "organization": "Mable",
            "responsibility": "",
            "source_text": "Provided individual support to clients through Mable, including appointments and assistance with daily living.\nMable | Jan 2020 – Present",
            "time_period_text": "Jan 2020 – Present",
        }
        with Session(self.engine) as session:
            resume = session.exec(select(Resume)).first()
            resume.source_text = source
            resume.experiences_json = json.dumps([experience])
            resume.ckb_json = '[{"sentinel":"unchanged"}]'
            safe_experiences = json.dumps([{
                "role_title": "Project Officer", "organization": "Example Agency",
                "responsibility": "Prepared reports and coordinated meetings.",
                "source_text": "Project Officer\nExample Agency\nFeb 2020 – Present\nPrepared reports and coordinated meetings.",
                "time_period_text": "Feb 2020 – Present",
            }])
            safe_resume = Resume(
                title="Safe Resume", source_text="Work Experience\n" + json.loads(safe_experiences)[0]["source_text"],
                experiences_json=safe_experiences, ckb_json='[{"safe":"unchanged"}]',
            )
            session.add_all([resume, safe_resume]); session.commit()

        scan = self.client.post("/resumes/risk-scan")

        self.assertEqual(scan.status_code, 200, scan.text)
        self.assertEqual(scan.json()["scanned_count"], 2)
        self.assertEqual(scan.json()["needs_review_count"], 1)
        flagged = scan.json()["resumes"][0]
        self.assertEqual(flagged["resume_id"], 1)
        self.assertEqual(flagged["experiences"][0]["index"], 1)
        self.assertEqual(
            flagged["experiences"][0]["review_reasons"],
            [
                "The role title looks like a description of duties rather than a job title.",
                "A possible job title immediately before this entry may have been left out.",
            ],
        )
        with Session(self.engine) as session:
            resume = session.exec(select(Resume).where(Resume.title != "Safe Resume")).first()
            marked = json.loads(resume.experiences_json)[0]
            self.assertEqual({key: marked[key] for key in experience}, experience)
            self.assertTrue(marked["needs_review"])
            self.assertEqual(resume.ckb_json, '[{"sentinel":"unchanged"}]')
            safe_resume = session.exec(select(Resume).where(Resume.title == "Safe Resume")).first()
            self.assertEqual(safe_resume.experiences_json, safe_experiences)
            self.assertEqual(safe_resume.ckb_json, '[{"safe":"unchanged"}]')

        second_scan = self.client.post("/resumes/risk-scan")
        self.assertEqual(second_scan.status_code, 200, second_scan.text)
        self.assertEqual(second_scan.json(), scan.json())
        with Session(self.engine) as session:
            second_marked_json = session.exec(select(Resume).where(Resume.title != "Safe Resume")).first().experiences_json
        self.assertEqual(second_marked_json, json.dumps([marked], ensure_ascii=False))

        blocked = self.client.post("/generate", json={"application_id": application_id, "document_type": "tailored_resume"})

        self.assertEqual(blocked.status_code, 409, blocked.text)
        self.assertEqual(blocked.json()["detail"]["code"], "master_resume_experience_needs_review")
        self.assertEqual(blocked.json()["detail"]["experiences"], flagged["experiences"])

    def test_real_nine_experience_backfill_changes_only_review_annotations_and_is_idempotent(self):
        original = json.loads(json.dumps(PRODUCTION_CLEAN_NINE_EXPERIENCES))
        with Session(self.engine) as session:
            session.add(Resume(
                title="Master Resume", source_text=PRODUCTION_MISSING_EXPERIENCES_SOURCE,
                experiences_json=json.dumps(original), ckb_json='[{"sentinel":"unchanged"}]',
            ))
            session.commit()

        first = self.client.post("/resumes/risk-scan")
        with Session(self.engine) as session:
            after_first = json.loads(session.exec(select(Resume)).first().experiences_json)
        second = self.client.post("/resumes/risk-scan")
        with Session(self.engine) as session:
            stored = session.exec(select(Resume)).first()
            after_second = json.loads(stored.experiences_json)

        self.assertEqual(first.status_code, 200, first.text)
        self.assertEqual(second.status_code, 200, second.text)
        self.assertEqual(first.json(), second.json())
        self.assertEqual(after_first, after_second)
        self.assertEqual(
            [{key: value for key, value in item.items() if key not in {"needs_review", "review_reasons"}} for item in after_second],
            original,
        )
        self.assertEqual([item["needs_review"] for item in after_second], [False, True, True, False, False, False, False, False, False])
        self.assertEqual(after_second[1]["review_reasons"], ["duty_shaped_role_title"])
        self.assertEqual(after_second[2]["review_reasons"], ["duty_shaped_role_title"])
        self.assertEqual(stored.ckb_json, '[{"sentinel":"unchanged"}]')

    def test_risky_resume_edit_returns_actionable_detail_without_changing_saved_resume(self):
        source = """Work Experience
Support Worker
Provided individual support to clients through Mable, including appointments and assistance with daily living.
Mable | Jan 2020 – Present
Education"""
        safe_source = "Work Experience\nProject Officer\nExample Agency\nFeb 2020 – Present\nPrepared reports and coordinated meetings."
        safe_experience = {
            "id": "saved", "role_title": "Project Officer", "organization": "Example Agency",
            "responsibility": "Prepared reports and coordinated meetings.",
            "source_text": "Project Officer\nExample Agency\nFeb 2020 – Present\nPrepared reports and coordinated meetings.",
            "time_period_text": "Feb 2020 – Present",
        }
        with Session(self.engine) as session:
            saved = Resume(title="Master Resume", source_text=safe_source,
                           experiences_json=json.dumps([safe_experience]), ckb_json="[]")
            session.add(saved); session.commit(); session.refresh(saved); resume_id = saved.id
        before = self.client.get("/resumes").json()[0]

        response = self.client.patch(f"/resumes/{resume_id}", json={
            "source_text": source,
            "experiences_json": json.dumps([{
                "id": "draft-edit", "role_title": "Provided individual support to clients through Mable, including appointments and assistance with daily living.",
                "organization": "Mable", "responsibility": "", "source_text": "Provided individual support to clients through Mable, including appointments and assistance with daily living.\nMable | Jan 2020 – Present",
                "time_period_text": "Jan 2020 – Present",
            }]),
        })

        self.assertEqual(response.status_code, 409, response.text)
        detail = response.json()["detail"]
        self.assertEqual(detail["code"], "master_resume_experience_needs_review")
        self.assertEqual(detail["experiences"][0]["id"], "draft-edit")
        self.assertNotIn("duty_shaped_role_title", str(detail))
        self.assertEqual(self.client.get("/resumes").json()[0], before)

    def test_production_shaped_historical_resume_requires_complete_reupload(self):
        roles = [
            ("Finance Administration Officer", "Department of Communities – Disability Services | WA State Government", "Feb 2026", "Aug 2026"),
            ("Executive Assistant to Board Member", "Avaintec", "Nov 2017", "Jan 2019"),
            ("Project Administration Officer", "China Communications Construction Company – Kenya Branch", "Jan 2016", "Aug 2017"),
            ("Project Administration Officer", "Chevron CDB Project", "Aug 2012", "Dec 2015"),
            ("Project Assistant", "Pratt & Whitney", "July 2007", "June 2012"),
        ]
        historical = [{
            "role_title": role, "organization": employer,
            "responsibility": f"Grounded responsibility for role {index} with sufficient detail.",
            "source_section": f"Work Experience > {employer} > {role}",
        } for index, (role, employer, _, _) in enumerate(roles)]
        historical[1]["responsibility"] = "Nov 2017 – Jan 2019 Prepared agendas and coordinated executive meetings."
        historical[2]["organization"] += " Jan 2016 – Aug 2017"
        stale = build_career_knowledge_base("Professional Summary\nAdministration\nCore Capabilities\nMicrosoft Office Suite", json.dumps(historical))

        with Session(self.engine) as session:
            resume = Resume(
                source_text="Professional Summary\nAdministration\nCore Capabilities\nMicrosoft Office Suite (Advanced Excel, Word, Outlook, Teams)",
                experiences_json=json.dumps(historical), ckb_json=json.dumps(stale),
            )
            application = JobApplication(
                company="Bennco", position_title="Project Administrator", job_description="Project administration",
                application_decision_json='{"status":"sentinel"}',
            )
            session.add_all([resume, application]); session.commit(); session.refresh(resume); session.refresh(application)
            application_id = application.id
            refreshed, status = get_or_refresh_current_ckb(session, resume, None)
            session.refresh(resume)

        self.assertEqual(status, "refreshed_stale")
        self.assertEqual(
            [item["time_period_status"] for item in refreshed],
            ["not_provided", "verified", "verified", "not_provided", "not_provided"],
        )
        self.assertIsNotNone(master_resume_integrity_issue(resume))
        self.assertNotIn("Nov 2017", json.loads(resume.experiences_json)[1]["responsibility"])
        self.assertNotIn("Jan 2016", json.loads(resume.experiences_json)[2]["organization"])
        for response in (
            self.client.post(f"/applications/{application_id}/decision"),
            self.client.post("/generate", json={"application_id": application_id, "document_type": "tailored_resume"}),
        ):
            self.assertEqual(response.status_code, 409)
            self.assertIn("Re-upload the complete Resume", response.json()["detail"])

        complete_source = "Professional Summary\nGrounded administration experience.\nCore Capabilities\nMicrosoft Office Suite\nWork Experience\n" + "\n".join(
            f"{role}\n{employer}\n{start} – {end}\nGrounded responsibility for {role} with enough detail."
            for role, employer, start, end in roles
        )
        document = Document()
        for line in complete_source.splitlines():
            document.add_paragraph(line)
        stream = BytesIO(); document.save(stream)
        uploaded = self.client.post(
            "/resumes/upload", files={"file": ("resume.docx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"title": "Master Resume"},
        )
        self.assertEqual(uploaded.status_code, 200, uploaded.text)
        with Session(self.engine) as session:
            reuploaded = session.exec(select(Resume).order_by(Resume.updated_at.desc())).first()
            canonical = json.loads(reuploaded.experiences_json)
            ckb = json.loads(reuploaded.ckb_json)
            application = session.get(JobApplication, application_id)
            self.assertEqual(application.application_decision_json, "{}")
            application.application_decision_json = '{"status":"sentinel"}'
            session.add(application); session.commit()
            resume_id = reuploaded.id
            persisted_experiences_json = reuploaded.experiences_json
            persisted_ckb_json = reuploaded.ckb_json
            self.assertIsNone(master_resume_integrity_issue(reuploaded))
        self.assertEqual(
            [item["time_period_text"] for item in canonical],
            [f"{start} – {end}" for _, _, start, end in roles],
        )
        self.assertEqual(
            [(item["time_period"]["start"], item["time_period"]["end"]) for item in ckb],
            [(start, end) for _, _, start, end in roles],
        )
        matches = {"matches": [{"criteria_id": "C1", "matched_evidence": [item["evidence_id"] for item in ckb], "match_type": "direct", "coverage": "strong"}]}
        plan = build_resume_curation_plan({"criteria": [{"criteria_id": "C1", "criteria_type": "essential"}]}, matches, ckb)
        self.assertEqual([item["display_period"] for item in plan["roles"]], [f"{start} - {end}" for _, _, start, end in roles])

        stale_patch = self.client.patch(f"/resumes/{resume_id}", json={
            "source_text": "Professional Summary\nAdministration\nCore Capabilities\nMicrosoft Office Suite",
            "experiences_json": json.dumps(canonical),
        })
        self.assertEqual(stale_patch.status_code, 409)
        self.assertIn("changed after this editor loaded", stale_patch.json()["detail"])
        with Session(self.engine) as session:
            unchanged = session.get(Resume, resume_id)
            self.assertEqual((unchanged.source_text, unchanged.experiences_json, unchanged.ckb_json), (complete_source, persisted_experiences_json, persisted_ckb_json))
            self.assertEqual(session.get(JobApplication, application_id).application_decision_json, '{"status":"sentinel"}')

        edited_source = complete_source.replace("Grounded administration experience.", "Grounded administration and project experience.")
        valid_patch = self.client.patch(f"/resumes/{resume_id}", json={
            "source_text": edited_source, "experiences_json": json.dumps(canonical),
        })
        self.assertEqual(valid_patch.status_code, 200, valid_patch.text)
        reloaded = self.client.get("/resumes").json()[0]
        unchanged_save = self.client.patch(f"/resumes/{resume_id}", json={
            "source_text": reloaded["source_text"], "experiences_json": reloaded["experiences_json"],
        })
        self.assertEqual(unchanged_save.status_code, 200, unchanged_save.text)
        with Session(self.engine) as session:
            application = session.get(JobApplication, application_id)
            requirements = self.requirements("resume")
            application.application_requirements_json = json.dumps(requirements)
            application.job_model_json = '{"schema_version":"1.0","criteria":[],"limit_scope":"unspecified"}'
            session.add(application); session.commit()
        with patch("app.main.match_evidence_batch", return_value={"schema_version": "1.0", "matches": [], "unused_evidence": []}):
            diagnosed = self.client.post(f"/applications/{application_id}/decision")
        self.assertEqual(diagnosed.status_code, 200, diagnosed.text)

    def test_bennco_stale_ckb_refreshes_once_and_enforces_all_five_periods(self):
        roles = [
            ("Finance Administration Officer", "Department of Communities – Disability Services", "Feb 2026", "Present"),
            ("Executive Assistant to Board Member", "Avaintec", "Nov 2017", "Jan 2019"),
            ("Project Administration Officer", "CCCC Kenya", "Jan 2016", "Aug 2017"),
            ("Project Administration Officer", "Chevron CDB Project", "Aug 2012", "Dec 2015"),
            ("Project Assistant", "Pratt & Whitney", "Oct 2007", "Aug 2012"),
        ]
        authoritative = [{
            "role_title": role, "organization": employer,
            "responsibility": f"Distinct grounded duty {index} involving verified records, systems, stakeholders, schedules and reporting support.",
            "source_section": f"Work Experience > {employer} > {role}",
            "source_text": f"{role}\n{employer}\n{start} – {end}\nDistinct grounded duty {index} involving verified records, systems, stakeholders, schedules and reporting support.",
        } for index, (role, employer, start, end) in enumerate(roles)]
        source_text = "Work Experience\n" + "\n".join(item["source_text"] for item in authoritative)
        experiences = json.loads(json.dumps(authoritative))
        for index, item in enumerate(experiences):
            if index in {1, 2}:
                item["time_period_text"] = f"{roles[index][2]} – {roles[index][3]}"
            else:
                item["source_text"] = f"{item['role_title']}\n{item['organization']}\n{item['responsibility']}"
        stale = build_career_knowledge_base(source_text, json.dumps(experiences))
        self.assertEqual(
            [item["time_period_status"] for item in stale],
            ["not_provided", "verified", "verified", "not_provided", "not_provided"],
        )

        with Session(self.engine) as session:
            resume = Resume(source_text=source_text, experiences_json=json.dumps(experiences), ckb_json=json.dumps(stale))
            application = JobApplication(company="Bennco", position_title="Office Administrator", job_description="Administration", application_decision_json='{"status":"ready"}')
            session.add_all([resume, application]); session.commit(); session.refresh(resume); session.refresh(application)
            refreshed, status = get_or_refresh_current_ckb(session, resume, None)
            self.assertEqual(application.application_decision_json, "{}")
            application.application_decision_json = '{"status":"sentinel"}'
            session.add(application); session.commit()
            reused, second_status = get_or_refresh_current_ckb(session, resume, None)
            session.refresh(application)

        self.assertEqual((status, second_status), ("refreshed_stale", "reused_current"))
        self.assertEqual(refreshed, reused)
        self.assertEqual(application.application_decision_json, '{"status":"sentinel"}')
        self.assertEqual(
            [(item["time_period"]["start"], item["time_period"]["end"], item["time_period_status"]) for item in refreshed],
            [(start, end, "verified") for _, _, start, end in roles],
        )
        matches = {"matches": [{"criteria_id": "C1", "matched_evidence": [item["evidence_id"] for item in refreshed], "match_type": "direct", "coverage": "strong"}]}
        plan = build_resume_curation_plan({"criteria": [{"criteria_id": "C1", "criteria_type": "essential"}]}, matches, refreshed)
        self.assertEqual([item["display_period"] for item in plan["roles"]], [f"{start} - {end}" for _, _, start, end in roles])

        headers = [f"### {role}\n{employer}" for role, employer, _, _ in roles]
        incomplete = "## Professional Summary\nGrounded.\n## Key Skills\nAdministration.\n## Work Experience\n" + "\n".join(
            f"{header}\n{start} – {end}" if index == 2 else header
            for index, (header, (_, _, start, end)) in enumerate(zip(headers, roles))
        )
        missing = validate_resume_content(incomplete, plan, [item["evidence_id"] for item in refreshed])
        self.assertEqual([item["code"] for item in missing["issues"]].count("missing_role_period"), 4)
        for missing_index in range(len(roles)):
            individually_incomplete = "## Professional Summary\nGrounded.\n## Key Skills\nAdministration.\n## Work Experience\n" + "\n".join(
                header if index == missing_index else f"{header}\n{start} – {end}"
                for index, (header, (_, _, start, end)) in enumerate(zip(headers, roles))
            )
            result = validate_resume_content(individually_incomplete, plan, [item["evidence_id"] for item in refreshed])
            self.assertIn("missing_role_period", [item["code"] for item in result["issues"]])
        complete = "## Professional Summary\nGrounded.\n## Key Skills\nAdministration.\n## Work Experience\n" + "\n".join(
            f"{header}\n{start} – {end}" for header, (_, _, start, end) in zip(headers, roles)
        )
        self.assertTrue(validate_resume_content(complete, plan, [item["evidence_id"] for item in refreshed])["valid"])

    def test_current_empty_date_states_are_reused_without_rebuild(self):
        for status in ("uncertain", "not_provided"):
            with self.subTest(status=status), Session(self.engine) as session:
                ckb = [{"schema_version": "2.0", "evidence_type": "experience", "source_group_id": "role", "time_period": {"start": None, "end": None}, "time_period_status": status}]
                resume = Resume(source_text="Authoritative source", experiences_json="[]", ckb_json=json.dumps(ckb))
                session.add(resume); session.commit(); session.refresh(resume)
                with patch("app.main.serialise_ckb", side_effect=AssertionError("current CKB must not rebuild")):
                    result, current_status = get_or_refresh_current_ckb(session, resume, None)
                self.assertEqual((result, current_status), (ckb, "reused_current"))

    def test_persisted_resume_integrity_blocks_release_despite_pack_and_ats_pass(self):
        application_id = self.seed(required=("resume",))
        roles = [
            ("Finance Administration Officer", "Department of Communities – Disability Services", "Feb 2026 – Present"),
            ("Executive Assistant to Board Member", "Avaintec", "Nov 2017 – Jan 2019"),
            ("Project Administration Officer", "China Communications Construction Company – Kenya Branch", "Jan 2016 – Aug 2017"),
            ("Project Administration Officer", "Chevron CDB Project", "Aug 2012 – Dec 2015"),
            ("Project Assistant", "Pratt & Whitney", "Oct 2007 – Aug 2012"),
        ]
        plan = {
            "schema_version": "1.1", "required_sections": ["Professional Summary", "Key Skills", "Work Experience"],
            "selected_evidence": [], "roles": [{
                "role_marker": role, "employer_marker": employer, "display_period": period,
                "chronology_order": index, "include_role_header": True,
            } for index, (role, employer, period) in enumerate(roles)],
        }
        prefix = "Alex Morgan\n0400000000 | alex@example.com\n## Professional Summary\nGrounded.\n## Key Skills\nAdministration.\n## Work Experience\n"
        incomplete = prefix + "\n".join(
            f"**{role}**\n{employer}" + (f"\n{period}" if index in {1, 2} else "")
            for index, (role, employer, period) in enumerate(roles)
        )
        complete = prefix + "\n".join(f"**{role}**\n{employer}\n{period}" for role, employer, period in roles)

        def persist(content):
            with Session(self.engine) as session:
                application = session.get(JobApplication, application_id)
                profile = session.exec(select(ApplicantProfile)).first()
                document = session.exec(select(GeneratedDocument).where(GeneratedDocument.application_id == application_id)).first()
                if not document:
                    document = GeneratedDocument(
                        application_id=application_id, document_type="tailored_resume", reviewer_json='{"status":"pass"}',
                        structured_content_json=json.dumps(plan), used_experiences_json="[]", content=content,
                    )
                    session.add(document); session.flush()
                else:
                    document.content = content
                pack = {"status": "pass", "blocks_release": False, "skipped": True, "skip_reason": "No comparison candidates.", "results": []}
                application.release_state_json = json.dumps({
                    "schema_version": "1.0",
                    "details_confirmation": {"fingerprint": details_fingerprint(application, profile)},
                    "pack_review": {"fingerprint": pack_fingerprint(application, profile, {"tailored_resume": document}), "result": pack},
                    "ats": {"document_id": document.id, "content_sha256": fingerprint(content), "format": "docx", "template": "classic", "result": {"status": "pass", "ready": True}},
                })
                session.add_all([application, document]); session.commit()

        persist(incomplete)
        final = self.client.get(f"/applications/{application_id}/quality-check").json()
        release = self.client.get(f"/applications/{application_id}/release-checklist").json()
        blocked = self.client.post(f"/applications/{application_id}/prepare-submission")

        self.assertEqual([item["code"] for item in final["issues"]].count("missing_role_period"), 3)
        self.assertFalse(final["ready"])
        self.assertTrue(release["checks"]["pack_review"]["ready"])
        self.assertTrue(release["checks"]["ats"]["ready"])
        self.assertFalse(release["ready"])
        self.assertNotEqual(release["status"], "ready_to_apply")
        self.assertEqual(blocked.status_code, 409)

        persist(complete)
        final = self.client.get(f"/applications/{application_id}/quality-check").json()
        release = self.client.get(f"/applications/{application_id}/release-checklist").json()
        self.assertTrue(final["ready"], final["issues"])
        self.assertTrue(release["ready"])
        self.assertEqual(release["status"], "ready_to_apply")
        self.assertEqual(self.client.post(f"/applications/{application_id}/prepare-submission").status_code, 200)

    def test_post_repair_resume_must_retain_authoritative_employment_block(self):
        plan = {
            "schema_version": "1.1", "required_sections": ["Professional Summary", "Key Skills", "Work Experience"],
            "selected_evidence": [], "roles": [{
                "employer_marker": "Avaintec", "role_marker": "Executive Assistant to Board Member",
                "display_period": "Nov 2017 - Jan 2019", "chronology_order": 0, "include_role_header": True,
            }],
        }
        valid = """## Professional Summary
Grounded support.
## Key Skills
Administration
## Work Experience
**Executive Assistant to Board Member**
**Avaintec**
Nov 2017 – Jan 2019"""
        missing = """## Professional Summary
Grounded support.
## Key Skills
Administration
## Work Experience
Other grounded work."""
        missing_date = valid.replace("Nov 2017 – Jan 2019", "")
        for repaired, expected_status, expected_error in (
            (valid, 200, ""),
            (missing, 502, "missing the required role header"),
            (missing_date, 502, "missing the authoritative employment period"),
        ):
            with self.subTest(expected_status=expected_status):
                application_id = self.seed(required=("resume",))
                with Session(self.engine) as session:
                    application = session.get(JobApplication, application_id)
                    profile = session.exec(select(ApplicantProfile)).first()
                    requirements = json.loads(application.application_requirements_json)
                    requirements["source"] = "source_aware_parser"
                    application.application_requirements_json = json.dumps(requirements)
                    application.application_decision_json = json.dumps({
                        "schema_version": "1.0", "status": "ready", "application_recommendation": "apply",
                        "inputs": decision_inputs(json.loads(application.job_model_json), requirements, [], profile),
                        "requirements": [], "questions": [], "blocking_issues": [],
                    })
                    session.add(application); session.commit()
                with patch("app.main.match_evidence_batch", return_value={"schema_version": "1.0", "matches": [], "unused_evidence": []}), patch(
                    "app.main.build_resume_curation_plan", return_value=plan
                ), patch("app.main.generate_draft", return_value=valid), patch(
                    "app.main.repair_tailored_resume", return_value=(repaired, {"status": "pass", "results": []})
                ):
                    response = self.client.post("/generate", json={"application_id": application_id, "document_type": "tailored_resume"})

                self.assertEqual(response.status_code, 200, response.text)
                if expected_status == 200:
                    self.assertEqual(json.loads(response.json()["trace_json"])["runtime"]["ckb_status"], "reused_current")
                if expected_status == 502:
                    review = json.loads(response.json()["reviewer_json"])
                    self.assertNotEqual(review["status"], "pass")
                    self.assertIn(expected_error, str(review["results"]))
                    self.assertTrue(response.json()["content"])

    def test_resume_name_is_restored_after_automatic_repair(self):
        application_id = self.seed(required=("resume",))
        with Session(self.engine) as session:
            application = session.get(JobApplication, application_id)
            profile = session.exec(select(ApplicantProfile)).first()
            requirements = json.loads(application.application_requirements_json)
            requirements["source"] = "source_aware_parser"
            application.application_requirements_json = json.dumps(requirements)
            application.application_decision_json = json.dumps({
                "schema_version": "1.0", "status": "ready", "application_recommendation": "apply",
                "inputs": decision_inputs(json.loads(application.job_model_json), requirements, [], profile),
                "requirements": [], "questions": [], "blocking_issues": [],
            })
            session.add(application); session.commit()
        draft = "## Professional Summary\nGrounded support.\n## Key Skills\nAdministration\n## Work Experience\nGrounded support."
        with patch("app.main.match_evidence_batch", return_value={"schema_version": "1.0", "matches": [], "unused_evidence": []}), patch(
            "app.main.generate_draft", return_value=draft
        ), patch("app.main.repair_tailored_resume", return_value=(draft, {"status": "pass", "results": []})):
            response = self.client.post("/generate", json={"application_id": application_id, "document_type": "tailored_resume"})

        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn("Alex Morgan", response.json()["content"])
        self.assertIn("0400000000", response.json()["content"])
        self.assertIn("alex@example.com", response.json()["content"])

    def test_reviewer_provider_failure_keeps_same_draft_for_safe_retry(self):
        application_id = self.seed(required=("cover_letter",))
        with Session(self.engine) as session:
            application = session.get(JobApplication, application_id)
            profile = session.exec(select(ApplicantProfile)).first()
            requirements = json.loads(application.application_requirements_json)
            requirements["source"] = "source_aware_parser"
            application.application_requirements_json = json.dumps(requirements)
            application.application_decision_json = json.dumps({
                "schema_version": "1.0", "status": "ready", "application_recommendation": "apply",
                "inputs": decision_inputs(json.loads(application.job_model_json), requirements, [], profile),
                "requirements": [], "questions": [], "blocking_issues": [],
            })
            session.add(application); session.commit()
        draft = "Alex Morgan\n0400000000 | alex@example.com\nApplication for Office Administrator\n\n" + "Grounded administration support. " * 40
        with patch("app.main.match_evidence_batch", return_value={"schema_version": "1.0", "matches": [], "unused_evidence": []}), patch(
            "app.main.generate_draft", return_value=draft
        ), patch("app.main.repair_cover_letter", side_effect=AIServiceError("Unterminated string at line 144")):
            failed = self.client.post("/generate", json={"application_id": application_id, "document_type": "cover_letter"})

        self.assertEqual(failed.status_code, 502, failed.text)
        self.assertNotIn("Unterminated", failed.json()["detail"]["message"])
        document_id = failed.json()["detail"]["document_id"]
        documents = self.client.get(f"/applications/{application_id}/documents").json()
        self.assertEqual([item["id"] for item in documents], [document_id])
        self.assertEqual(json.loads(documents[0]["reviewer_json"])["status"], "provider_failed")

        with patch("app.main.review_cover_letter", return_value={"status": "pass", "results": []}):
            retried = self.client.post(f"/documents/{document_id}/review")
        self.assertEqual(retried.status_code, 200)
        self.assertEqual(retried.json()["id"], document_id)
        self.assertEqual(json.loads(retried.json()["reviewer_json"])["status"], "pass")

    def setUp(self):
        self.engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        SQLModel.metadata.create_all(self.engine)

        def sessions():
            with Session(self.engine) as session:
                yield session

        app.dependency_overrides[get_session] = sessions
        self.client = TestClient(app)

    def tearDown(self):
        app.dependency_overrides.clear()
        self.engine.dispose()

    @staticmethod
    def requirements(*required):
        result = empty_application_requirements("User-confirmed requirements", "user_supplied")
        result["review_status"] = "confirmed"
        for document in result["documents"].values():
            document.update(requirement="not_required", format="not_applicable")
        for name in required:
            result["documents"][name].update(requirement="required", format="standalone")
        return result

    def seed(self, required=("resume", "cover_letter")):
        with Session(self.engine) as session:
            profile = ApplicantProfile(first_name="Alex", last_name="Morgan", phone="0400000000", email="alex@example.com")
            resume = Resume(source_text="Alex Morgan\nOffice administration experience.", ckb_json="[]")
            requirements = self.requirements(*required)
            application = JobApplication(
                company="SRG Global", position_title="Office Administrator", job_url="https://example.com/apply",
                job_description="Office administration and procurement support.",
                application_requirements_json=json.dumps(requirements), job_model_json='{"schema_version":"1.0","criteria":[],"limit_scope":"unspecified"}',
            )
            session.add_all([profile, resume, application]); session.flush()
            application.application_decision_json = json.dumps({
                "schema_version": "1.0", "status": "ready", "application_recommendation": "apply",
                "inputs": decision_inputs(json.loads(application.job_model_json), requirements, [], profile),
                "requirements": [], "questions": [], "blocking_issues": [],
            })
            session.add(application); session.commit(); session.refresh(application)
            return application.id

    def test_not_required_selection_criteria_is_rejected_and_cannot_block_final_check(self):
        application_id = self.seed()
        with Session(self.engine) as session:
            application = session.get(JobApplication, application_id)
            application.selection_plan_json = '{"items":[{"criteria_id":"C1","evidence_status":"weak"}]}'
            session.add_all([
                GeneratedDocument(application_id=application_id, document_type="tailored_resume", content="Alex Morgan\n## Professional Summary\nAdmin\n## Key Skills\nAdmin\n## Work Experience\nAdmin", reviewer_json='{"status":"pass"}', structured_content_json='{"schema_version":"1.1","roles":[]}'),
                GeneratedDocument(application_id=application_id, document_type="cover_letter", content="Application for Office Administrator at SRG Global. " + "evidence " * 220, reviewer_json='{"status":"pass"}'),
                GeneratedDocument(application_id=application_id, document_type="selection_criteria", content="Unnecessary", reviewer_json='{"status":"fail","generation_status":"needs_ckb_update","results":[]}'),
            ])
            session.commit()

        rejected = self.client.post("/generate", json={"application_id": application_id, "document_type": "selection_criteria"})
        final = self.client.get(f"/applications/{application_id}/quality-check").json()
        release = self.client.get(f"/applications/{application_id}/release-checklist").json()
        exported = self.client.get(f"/applications/{application_id}/export-pack?format=docx&template=classic")

        self.assertEqual(rejected.status_code, 409)
        self.assertNotIn("selection_criteria", final["checked_documents"])
        self.assertFalse(any(item.get("document_type") == "selection_criteria" for item in final["issues"]))
        self.assertTrue(release["checks"]["selection_confirmations"]["ready"])
        self.assertNotIn("selection_criteria", [item["document_type"] for item in self.client.get(f"/applications/{application_id}/documents").json()])
        self.assertEqual(exported.status_code, 200)
        with ZipFile(BytesIO(exported.content)) as archive:
            self.assertFalse(any("Selection_Criteria" in name for name in archive.namelist()))

    def test_job_edit_updates_primary_source_and_makes_old_documents_historical(self):
        application_id = self.seed()
        with Session(self.engine) as session:
            application = session.get(JobApplication, application_id)
            session.add(JobSource(
                application_id=application_id, source_id="primary", source_type="primary_advertisement",
                title="Job advertisement", label="Primary advertisement", acquisition_status="fetched",
                extraction_status="extracted", classification_confidence="high", extracted_text=application.job_description,
            ))
            session.add(GeneratedDocument(application_id=application_id, document_type="tailored_resume", content="Old CV", reviewer_json='{"status":"pass"}'))
            session.commit()

        updated = self.client.patch(f"/applications/{application_id}", json={"job_description": "Changed procurement and workforce administration duties."})
        with Session(self.engine) as session:
            application = session.get(JobApplication, application_id)
            profile = session.exec(select(ApplicantProfile)).first()
            requirements = self.requirements("resume", "cover_letter")
            application.application_requirements_json = json.dumps(requirements)
            application.application_decision_json = json.dumps({
                "schema_version": "1.0", "status": "ready", "application_recommendation": "apply",
                "inputs": decision_inputs(json.loads(application.job_model_json), requirements, [], profile),
                "requirements": [], "questions": [], "blocking_issues": [],
            })
            session.commit()
        documents = self.client.get(f"/applications/{application_id}/documents").json()
        final = self.client.get(f"/applications/{application_id}/quality-check").json()
        with Session(self.engine) as session:
            source = session.exec(select(JobSource).where(JobSource.application_id == application_id)).first()

        self.assertEqual(updated.status_code, 200)
        self.assertEqual(documents, [])
        self.assertIn("Changed procurement", source.extracted_text)
        self.assertIn("stale_generated_document", [item["code"] for item in final["issues"]])
        with Session(self.engine) as session:
            application = session.get(JobApplication, application_id)
            resume = session.exec(select(Resume)).first()
            profile = session.exec(select(ApplicantProfile)).first()
            session.add(GeneratedDocument(
                application_id=application_id, document_type="tailored_resume", content="Current CV",
                reviewer_json='{"status":"pass"}',
                trace_json=json.dumps({"input_fingerprint": generation_inputs_fingerprint(application, resume, profile)}),
            ))
            session.commit()
            self.assertEqual(len(session.exec(select(GeneratedDocument).where(GeneratedDocument.application_id == application_id)).all()), 2)
        self.assertEqual(len(self.client.get(f"/applications/{application_id}/documents").json()), 1)

    def test_manual_cover_letter_edit_can_be_re_reviewed(self):
        application_id = self.seed(required=("cover_letter",))
        with Session(self.engine) as session:
            document = GeneratedDocument(
                application_id=application_id, document_type="cover_letter", content="Alex Morgan\n0400000000 | alex@example.com\nEdited grounded letter.",
                reviewer_json="{}", structured_content_json='{"priorities":["administration"]}',
            )
            session.add(document); session.commit(); session.refresh(document); document_id = document.id

        aggregate = {"allowed_claim": "12+ years of total employment experience"}
        with patch("app.main.aggregate_experience", return_value=aggregate) as calculate, \
             patch("app.main.review_cover_letter", return_value={"status": "pass", "results": []}) as reviewer:
            response = self.client.post(f"/documents/{document_id}/review")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(json.loads(response.json()["reviewer_json"])["status"], "pass")
        calculate.assert_called_once_with([])
        self.assertIn(aggregate["allowed_claim"], reviewer.call_args.args[3])
        self.assertEqual(json.loads(response.json()["trace_json"])["aggregate_experience"], aggregate)

    def test_failed_tailored_resume_can_be_re_reviewed_without_regeneration(self):
        application_id = self.seed()
        with Session(self.engine) as session:
            document = GeneratedDocument(
                application_id=application_id, document_type="tailored_resume",
                content="Alex Morgan\n0400000000 | alex@example.com\n## Work Experience\nGrounded administration.",
                reviewer_json='{"status":"fail","results":[{"issues":[{"type":"evidence_mismatch","description":"Old false positive"}]}]}',
                structured_content_json='{"roles":[]}',
            )
            session.add(document); session.commit(); session.refresh(document); document_id = document.id

        aggregate = {"allowed_claim": "12+ years of total employment experience"}
        with patch("app.main.aggregate_experience", return_value=aggregate) as calculate, \
             patch("app.main.review_tailored_resume", return_value={"status": "pass", "results": []}) as reviewer:
            response = self.client.post(f"/documents/{document_id}/review")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["id"], document_id)
        self.assertEqual(json.loads(response.json()["reviewer_json"])["status"], "pass")
        calculate.assert_called_once_with([])
        self.assertIn(aggregate["allowed_claim"], reviewer.call_args.args[4])
        self.assertEqual(json.loads(response.json()["trace_json"])["aggregate_experience"], aggregate)

    def test_manual_selection_criteria_edit_preserves_plan_for_re_review(self):
        application_id = self.seed(required=("selection_criteria",))
        plan = {"items": [{"criteria_id": "C1", "criteria_text": "Communication", "matched_evidence": [], "allocated_word_limit": 100}]}
        bundle = {
            "selection_plan": plan,
            "responses": [{"criteria_id": "C1", "final_response": "Original.", "evidence_used": [], "star": {"situation": "", "task": "", "action": "", "result": ""}}],
        }
        with Session(self.engine) as session:
            document = GeneratedDocument(
                application_id=application_id, document_type="selection_criteria",
                content="## Communication\n\nEdited grounded response.", reviewer_json="{}",
                structured_content_json=json.dumps(bundle),
            )
            session.add(document); session.commit(); session.refresh(document); document_id = document.id

        with patch("app.main.review_selection_criteria_batch", return_value={"status": "pass", "results": []}):
            response = self.client.post(f"/documents/{document_id}/review")

        self.assertEqual(response.status_code, 200)
        structured = json.loads(response.json()["structured_content_json"])
        self.assertEqual(structured["responses"][0]["final_response"], "Edited grounded response.")

    def test_submission_snapshot_excludes_unrequired_selection_criteria(self):
        application = JobApplication(company="SRG Global", position_title="Office Administrator", job_description="JD")
        documents = [
            GeneratedDocument(id=1, application_id=1, document_type="tailored_resume", content="CV"),
            GeneratedDocument(id=2, application_id=1, document_type="cover_letter", content="Letter"),
            GeneratedDocument(id=3, application_id=1, document_type="selection_criteria", content="Historical SC"),
        ]
        snapshot = build_submission_snapshot(application, documents, [], "Australia", [], documents[0].created_at, ("tailored_resume", "cover_letter"))
        self.assertNotIn("selection_criteria", snapshot["documents"])


if __name__ == "__main__":
    unittest.main()
